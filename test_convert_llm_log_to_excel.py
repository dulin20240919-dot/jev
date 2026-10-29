import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from convert_llm_log_to_excel import (
    COLUMNS,
    EXCEL_CELL_LIMIT,
    canonical_model_output,
    convert,
    extract_noul,
    record_to_row,
    write_excel,
)


def sample_record(model_output: str | None = None, noul: float = 0.2) -> dict:
    domain_llm = (
        "turn3.task1:TimeManage(target=Alarm,function=Create,"
        "value={time=每天早上7\\:45, event=早睡早起，不想歪事儿})"
    )
    tool_llm = (
        "turn4.task1:CreateAlarm(event=早睡早起，不想歪事儿,time=每天早上7\\:45)"
    )
    if model_output is None:
        model_output = domain_llm
    inner = {
        "model": "jev-latest",
        "answers": {
            "verify_tool_call": {
                "type": "noul",
                "noul": noul,
                "pred_index": 0,
                "probabilities": {"false": 0.8, "true": 0.9},
            }
        },
    }
    inner_text = json.dumps(inner, ensure_ascii=False)
    jev_prompt = {
        "model": "jev-latest-9b",
        "state": {
            "tools": [
                {"name": "TimeManage", "description": "时间管理的统一操作工具"}
            ],
            "other_candidates": [
                {"name": "NavigateControl", "description": "导航规划与管理"},
                {"name": "PhoneCall", "description": "通讯电话垂域工具"},
            ],
            "model_output": model_output,
            "messages": [
                {"role": "user", "content": "以后每天提醒我早睡早起，不想歪事儿"},
                {"role": "user", "content": "上一个15分钟的闹钟"},
            ],
        },
    }
    return {
        "asrText": "上一个15分钟的闹钟",
        "intent": "CHAT",
        "speakText": "",
        "time": "1793243918539",
        "sessionId": "5f3106fd-f6fd-4caa-ae73-373bf6645370",
        "interactionId": "4",
        "llm": [
            {
                "promptId": "DOMAIN_SELECT_TURBO_PY_PRO",
                "prompt": "<|im_start|>system\n# 角色定义",
                "llmText": domain_llm,
            },
            {
                "promptId": "TOOL_AND_EVENT_SELECT_PY_PRO",
                "prompt": "<|im_start|>system\n你是一个AI助手系统中的工具调用助手",
                "llmText": tool_llm,
            },
            {
                "promptId": "jev",
                "prompt": json.dumps(jev_prompt, ensure_ascii=False),
                "llmText": json.dumps(
                    {
                        "code": 0,
                        "modelRequestInfo": {
                            "contentBean": {"oriResponseBody": inner_text}
                        },
                        "oriResponseBody": inner_text,
                    },
                    ensure_ascii=False,
                ),
            },
        ],
    }


class ConvertTests(unittest.TestCase):
    def test_domain_model_is_selected_when_output_matches_its_llmtext(self):
        row = record_to_row(sample_record())
        self.assertEqual(row["asrText"], "上一个15分钟的闹钟")
        self.assertEqual(row["intent"], "CHAT")
        self.assertEqual(row["sessionId"], "5f3106fd-f6fd-4caa-ae73-373bf6645370")
        self.assertEqual(row["interactionId"], "4")
        self.assertEqual(row["time"], "1793243918539")
        self.assertEqual(row["jev判断模型"], "DOMAIN_SELECT_TURBO_PY_PRO")
        self.assertIn("角色定义", row["DOMAIN_SELECT_TURBO_PY_PRO的prompt"])
        self.assertIn("TimeManage", row["DOMAIN_SELECT_TURBO_PY_PRO的llmText"])
        self.assertIn("工具调用助手", row["TOOL_AND_EVENT_SELECT_PY_PRO的prompt"])
        self.assertIn("CreateAlarm", row["TOOL_AND_EVENT_SELECT_PY_PRO的llmText"])
        self.assertAlmostEqual(row["noul"], 0.2)
        messages = json.loads(row["message"])
        self.assertEqual(messages[-1]["content"], "上一个15分钟的闹钟")
        tools = json.loads(row["tools"])
        self.assertEqual(tools[0]["name"], "TimeManage")
        others = json.loads(row["other_candidates"])
        self.assertEqual(others[0]["name"], "NavigateControl")

    def test_tool_model_is_selected_when_output_matches_its_llmtext(self):
        tool_llm = sample_record()["llm"][1]["llmText"]
        row = record_to_row(sample_record(model_output=tool_llm))
        self.assertEqual(row["jev判断模型"], "TOOL_AND_EVENT_SELECT_PY_PRO")

    def test_extra_backslash_before_colon_still_matches(self):
        domain_llm = sample_record()["llm"][0]["llmText"]
        doubled = domain_llm.replace("\\:", "\\\\:")
        self.assertNotEqual(doubled, domain_llm)
        self.assertEqual(canonical_model_output(doubled), canonical_model_output(domain_llm))
        row = record_to_row(sample_record(model_output=doubled))
        self.assertEqual(row["jev判断模型"], "DOMAIN_SELECT_TURBO_PY_PRO")

    def test_unmatched_output_is_marked(self):
        row = record_to_row(sample_record(model_output="完全对不上的调用"))
        self.assertEqual(row["jev判断模型"], "未匹配")

    def test_noul_is_not_taken_from_probabilities(self):
        record = sample_record(noul=0.009708476135311471)
        self.assertAlmostEqual(extract_noul(record["llm"][2]["llmText"]), 0.009708476135311471)

    def test_noul_regex_fallback_when_body_is_not_json(self):
        text = (
            'prefix "verify_tool_call": {"type": "noul", "noul": 0.25, '
            '"probabilities": {"true": 0.1}} tail'
        )
        self.assertAlmostEqual(extract_noul(text), 0.25)

    def test_last_duplicate_prompt_id_wins(self):
        record = sample_record()
        record["llm"].append(
            {
                "promptId": "DOMAIN_SELECT_TURBO_PY_PRO",
                "prompt": "第二次领域选择",
                "llmText": record["llm"][0]["llmText"],
            }
        )
        row = record_to_row(record)
        self.assertEqual(row["DOMAIN_SELECT_TURBO_PY_PRO的prompt"], "第二次领域选择")

    def test_missing_jev_leaves_judgement_blank(self):
        record = sample_record()
        record["llm"] = record["llm"][:2]
        row = record_to_row(record)
        self.assertEqual(row["jev判断模型"], "")
        self.assertEqual(row["message"], "")
        self.assertIsNone(row["noul"])

    def test_excel_roundtrip_preserves_columns(self):
        row = record_to_row(sample_record())
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "out.xlsx"
            write_excel([row], path)
            sheet = load_workbook(path).active
            headers = [sheet.cell(1, col).value for col in range(1, len(COLUMNS) + 1)]
            self.assertEqual(headers, COLUMNS)
            values = {
                headers[col - 1]: sheet.cell(2, col).value
                for col in range(1, len(COLUMNS) + 1)
            }
            self.assertEqual(values["asrText"], "上一个15分钟的闹钟")
            self.assertEqual(values["jev判断模型"], "DOMAIN_SELECT_TURBO_PY_PRO")
            self.assertAlmostEqual(values["noul"], 0.2)
            self.assertIn("上一个15分钟的闹钟", values["message"])
            self.assertEqual(sheet.freeze_panes, "A2")

    def test_long_cell_is_truncated_to_excel_limit(self):
        row = record_to_row(sample_record())
        row["jev的prompt"] = "甲" * (EXCEL_CELL_LIMIT + 50)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "long.xlsx"
            write_excel([row], path)
            sheet = load_workbook(path).active
            prompt_col = COLUMNS.index("jev的prompt") + 1
            text = sheet.cell(2, prompt_col).value
            self.assertLessEqual(len(text), EXCEL_CELL_LIMIT)
            self.assertTrue(text.endswith("[truncated]"))

    def test_record_with_list_field_is_not_unwrapped(self):
        record = sample_record()
        record["list"] = ["不是另一条日志"]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "one.json"
            path.write_text(json.dumps(record, ensure_ascii=False), encoding="utf-8")
            rows = convert([path])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["asrText"], "上一个15分钟的闹钟")

    def test_json_array_jsonl_and_directory(self):
        first = sample_record()
        second = sample_record(
            model_output=first["llm"][1]["llmText"],
        )
        second["asrText"] = "再定一个晚上的"
        second["interactionId"] = "5"
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            array_path = root / "batch.json"
            array_path.write_text(
                json.dumps({"records": [first]}, ensure_ascii=False),
                encoding="utf-8",
            )
            jsonl_path = root / "more.jsonl"
            jsonl_path.write_text(
                json.dumps(second, ensure_ascii=False) + "\n",
                encoding="utf-8",
            )
            rows = convert([root])
            self.assertEqual(len(rows), 2)
            self.assertEqual(rows[0]["asrText"], "上一个15分钟的闹钟")
            self.assertEqual(rows[0]["jev判断模型"], "DOMAIN_SELECT_TURBO_PY_PRO")
            self.assertEqual(rows[1]["asrText"], "再定一个晚上的")
            self.assertEqual(rows[1]["jev判断模型"], "TOOL_AND_EVENT_SELECT_PY_PRO")


if __name__ == "__main__":
    unittest.main()
