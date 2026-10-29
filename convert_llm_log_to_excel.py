#!/usr/bin/env python3
"""把语音助手 LLM 日志转成 Excel。

每条记录一行，列包括 asrText、两段上游模型的 prompt / llmText、
jev 正在判别的模型、jev 的 prompt 及其 state 里的 message / tools /
other_candidates、jev 的 llmText，以及 verify_tool_call 的 noul。

依赖: pip install openpyxl

用法:
    python convert_llm_log_to_excel.py record.json -o record.xlsx
    python convert_llm_log_to_excel.py records.jsonl -o records.xlsx
    python convert_llm_log_to_excel.py logs/ -o logs.xlsx

输入可以是单个 JSON 对象、JSON 数组，或 JSONL（一行一条）。
目录会按文件名顺序读取其中的 .json / .jsonl / .ndjson。
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DOMAIN_PROMPT_ID = "DOMAIN_SELECT_TURBO_PY_PRO"
TOOL_PROMPT_ID = "TOOL_AND_EVENT_SELECT_PY_PRO"
JEV_PROMPT_ID = "jev"

# Excel 单元格上限是 32767。截断时留出标记，避免写入失败。
EXCEL_CELL_LIMIT = 32767
TRUNCATION_MARK = "\n...[truncated]"

COLUMNS = [
    "asrText",
    "intent",
    "sessionId",
    "interactionId",
    "time",
    "DOMAIN_SELECT_TURBO_PY_PRO的prompt",
    "DOMAIN_SELECT_TURBO_PY_PRO的llmText",
    "TOOL_AND_EVENT_SELECT_PY_PRO的prompt",
    "TOOL_AND_EVENT_SELECT_PY_PRO的llmText",
    "jev判断模型",
    "jev的prompt",
    "message",
    "tools",
    "other_candidates",
    "jev的llmText",
    "noul",
]

COLUMN_WIDTHS = {
    "asrText": 28,
    "intent": 14,
    "sessionId": 40,
    "interactionId": 16,
    "time": 18,
    "DOMAIN_SELECT_TURBO_PY_PRO的prompt": 48,
    "DOMAIN_SELECT_TURBO_PY_PRO的llmText": 56,
    "TOOL_AND_EVENT_SELECT_PY_PRO的prompt": 48,
    "TOOL_AND_EVENT_SELECT_PY_PRO的llmText": 56,
    "jev判断模型": 36,
    "jev的prompt": 48,
    "message": 64,
    "tools": 42,
    "other_candidates": 42,
    "jev的llmText": 48,
    "noul": 24,
}

# openpyxl 拒绝写入的 XML 1.0 非法控制字符。
_ILLEGAL_XML_CHARS = re.compile(r"[\000-\010\013\014\016-\037]")
_NOUL_LITERAL = re.compile(
    r'"verify_tool_call"\s*:\s*\{[^{}]*?"noul"\s*:\s*(-?\d+(?:\.\d+)?(?:[eE][+-]?\d+)?)',
    re.DOTALL,
)
_INPUT_SUFFIXES = {".json", ".jsonl", ".ndjson"}


def parse_json_maybe(text: str) -> Any:
    """解析 JSON 文本。前面若有非 JSON 前缀，从第一个对象或数组开始解析。"""
    stripped = text.strip()
    if not stripped:
        return None
    try:
        return json.loads(stripped)
    except json.JSONDecodeError:
        pass
    start = -1
    for index, char in enumerate(stripped):
        if char in "{[":
            start = index
            break
    if start < 0:
        return None
    try:
        value, _ = json.JSONDecoder().raw_decode(stripped[start:])
    except json.JSONDecodeError:
        return None
    return value


def dump_json(value: Any) -> str:
    """把对象写成可读的 JSON。字符串原样保留。"""
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    return json.dumps(value, ensure_ascii=False, indent=2)


def canonical_model_output(text: str) -> str:
    """抹平多层 JSON 转义造成的反斜杠差异。

    日志里的时间常写成 ``7\\:45``。jev prompt 再套一层字符串后，
    同一段文本可能多出成对的反斜杠。比较前把连续的 ``\\\\:`` 收成 ``\\:``。
    """
    current = text.strip()
    while True:
        collapsed = current.replace("\\\\:", "\\:")
        if collapsed == current:
            return current
        current = collapsed


def index_llm(llm_items: Any) -> dict[str, dict[str, Any]]:
    """按 promptId 索引。同一 promptId 出现多次时保留最后一次。"""
    indexed: dict[str, dict[str, Any]] = {}
    if not isinstance(llm_items, list):
        return indexed
    for item in llm_items:
        if not isinstance(item, dict):
            continue
        prompt_id = str(item.get("promptId", "")).strip()
        if prompt_id:
            indexed[prompt_id] = item
    return indexed


def field_text(item: dict[str, Any] | None, key: str) -> str:
    if not item:
        return ""
    return dump_json(item.get(key))


def jev_state(prompt_text: str) -> dict[str, Any]:
    parsed = parse_json_maybe(prompt_text) if prompt_text else None
    if not isinstance(parsed, dict):
        return {}
    state = parsed.get("state")
    if isinstance(state, str):
        decoded = parse_json_maybe(state)
        if isinstance(decoded, dict):
            state = decoded
    if isinstance(state, dict):
        return state
    return {}


def judge_model(model_output: Any, domain_llm: str, tool_llm: str) -> str:
    """model_output 与哪一段 llmText 相同，jev 就在判别哪个模型。"""
    if model_output is None:
        return ""
    output_text = model_output if isinstance(model_output, str) else dump_json(model_output)
    if not output_text.strip():
        return ""

    candidates = [
        (DOMAIN_PROMPT_ID, domain_llm),
        (TOOL_PROMPT_ID, tool_llm),
    ]
    exact = [
        name
        for name, text in candidates
        if text.strip() and text.strip() == output_text.strip()
    ]
    if exact:
        return " | ".join(exact)

    output_canonical = canonical_model_output(output_text)
    normalized = [
        name
        for name, text in candidates
        if text.strip() and canonical_model_output(text) == output_canonical
    ]
    if normalized:
        return " | ".join(normalized)
    return "未匹配"


def _noul_from_answers(payload: Any) -> float | None:
    if not isinstance(payload, dict):
        return None
    answers = payload.get("answers")
    if not isinstance(answers, dict):
        return None
    verdict = answers.get("verify_tool_call")
    if not isinstance(verdict, dict) or "noul" not in verdict:
        return None
    value = verdict["noul"]
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return float(value)


def _walk_for_noul(payload: Any, depth: int = 0) -> float | None:
    if payload is None or depth > 8:
        return None
    found = _noul_from_answers(payload)
    if found is not None:
        return found
    if isinstance(payload, dict):
        for key, value in payload.items():
            if isinstance(value, str) and ("noul" in value or key == "oriResponseBody"):
                decoded = parse_json_maybe(value)
                if decoded is not None and not isinstance(decoded, str):
                    found = _walk_for_noul(decoded, depth + 1)
                    if found is not None:
                        return found
            found = _walk_for_noul(value, depth + 1)
            if found is not None:
                return found
    elif isinstance(payload, list):
        for item in payload:
            found = _walk_for_noul(item, depth + 1)
            if found is not None:
                return found
    return None


def extract_noul(llm_text: str) -> float | None:
    """从 jev llmText 中取出 answers.verify_tool_call.noul。

    llmText 外层是网关包装，真正的 answers 通常在 oriResponseBody 字符串里，
    而且可能再嵌一层。优先走 JSON 路径，解析失败时用原文里的数字字面量。
    """
    if not llm_text:
        return None
    parsed = parse_json_maybe(llm_text)
    found = _walk_for_noul(parsed)
    if found is not None:
        return found
    match = _NOUL_LITERAL.search(llm_text)
    if not match:
        return None
    return float(match.group(1))


def record_to_row(record: dict[str, Any]) -> dict[str, Any]:
    llm_by_id = index_llm(record.get("llm"))
    domain = llm_by_id.get(DOMAIN_PROMPT_ID)
    tool = llm_by_id.get(TOOL_PROMPT_ID)
    jev = llm_by_id.get(JEV_PROMPT_ID)

    domain_llm = field_text(domain, "llmText")
    tool_llm = field_text(tool, "llmText")
    jev_prompt = field_text(jev, "prompt")
    jev_llm = field_text(jev, "llmText")
    state = jev_state(jev_prompt)
    judged = judge_model(state.get("model_output"), domain_llm, tool_llm)
    if jev and not judged:
        judged = "未匹配"

    row: dict[str, Any] = {
        "asrText": dump_json(record.get("asrText")),
        "intent": dump_json(record.get("intent")),
        "sessionId": dump_json(record.get("sessionId")),
        "interactionId": dump_json(record.get("interactionId")),
        "time": dump_json(record.get("time")),
        "DOMAIN_SELECT_TURBO_PY_PRO的prompt": field_text(domain, "prompt"),
        "DOMAIN_SELECT_TURBO_PY_PRO的llmText": domain_llm,
        "TOOL_AND_EVENT_SELECT_PY_PRO的prompt": field_text(tool, "prompt"),
        "TOOL_AND_EVENT_SELECT_PY_PRO的llmText": tool_llm,
        "jev判断模型": judged,
        "jev的prompt": jev_prompt,
        "message": dump_json(state.get("messages")) if "messages" in state else "",
        "tools": dump_json(state.get("tools")) if "tools" in state else "",
        "other_candidates": (
            dump_json(state.get("other_candidates"))
            if "other_candidates" in state
            else ""
        ),
        "jev的llmText": jev_llm,
        "noul": extract_noul(jev_llm),
    }
    return row


def excel_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, float):
        return value
    text = value if isinstance(value, str) else str(value)
    text = _ILLEGAL_XML_CHARS.sub("", text)
    if len(text) <= EXCEL_CELL_LIMIT:
        return text
    keep = EXCEL_CELL_LIMIT - len(TRUNCATION_MARK)
    return text[:keep] + TRUNCATION_MARK


def write_excel(rows: list[dict[str, Any]], output: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "llm_log"

    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_alignment = Alignment(vertical="top", wrap_text=True)
    thin = Border(
        left=Side(style="thin", color="D9E2F3"),
        right=Side(style="thin", color="D9E2F3"),
        top=Side(style="thin", color="D9E2F3"),
        bottom=Side(style="thin", color="D9E2F3"),
    )
    alt_fill = PatternFill("solid", fgColor="F7F9FC")

    for col_index, name in enumerate(COLUMNS, start=1):
        cell = sheet.cell(1, col_index, name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin
        sheet.column_dimensions[get_column_letter(col_index)].width = COLUMN_WIDTHS[name]

    for row_index, row in enumerate(rows, start=2):
        for col_index, name in enumerate(COLUMNS, start=1):
            cell = sheet.cell(row_index, col_index, excel_cell(row.get(name)))
            cell.alignment = cell_alignment
            cell.border = thin
            if row_index % 2 == 0:
                cell.fill = alt_fill
            if name == "noul" and isinstance(row.get(name), float):
                cell.number_format = "0.################"
        sheet.row_dimensions[row_index].height = 72

    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = "A2"
    last_row = max(len(rows) + 1, 1)
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last_row}"
    sheet.sheet_view.showGridLines = False

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def _records_from_json(data: Any, source: Path) -> list[dict[str, Any]]:
    if isinstance(data, list):
        records = data
    elif isinstance(data, dict):
        # 带 llm / asrText 的对象就是一条日志，即使里面还有别的数组字段。
        if "llm" in data or "asrText" in data:
            records = [data]
        else:
            records = None
            for key in ("data", "records", "items", "logs", "list"):
                nested = data.get(key)
                if isinstance(nested, list):
                    records = nested
                    break
            if records is None:
                records = [data]
    else:
        raise ValueError(f"{source} 的顶层不是 JSON 对象或数组")

    normalized: list[dict[str, Any]] = []
    for index, record in enumerate(records, start=1):
        if not isinstance(record, dict):
            raise ValueError(f"{source} 第 {index} 条不是 JSON 对象")
        normalized.append(record)
    return normalized


def load_records(path: Path) -> list[dict[str, Any]]:
    text = path.read_text(encoding="utf-8-sig")
    suffix = path.suffix.lower()
    if suffix in {".jsonl", ".ndjson"}:
        records: list[dict[str, Any]] = []
        for line_no, line in enumerate(text.splitlines(), start=1):
            if not line.strip():
                continue
            try:
                item = json.loads(line)
            except json.JSONDecodeError as exc:
                raise ValueError(f"{path}:{line_no} 不是合法 JSON ({exc})") from exc
            if not isinstance(item, dict):
                raise ValueError(f"{path}:{line_no} 不是 JSON 对象")
            records.append(item)
        return records

    try:
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        raise ValueError(f"{path} 不是合法 JSON ({exc})") from exc
    return _records_from_json(data, path)


def iter_input_files(inputs: Iterable[Path]) -> list[Path]:
    files: list[Path] = []
    for raw in inputs:
        path = raw.expanduser()
        if path.is_dir():
            found = [
                item
                for item in path.rglob("*")
                if item.is_file() and item.suffix.lower() in _INPUT_SUFFIXES
            ]
            if not found:
                raise ValueError(f"目录里没有 JSON / JSONL 文件: {path}")
            files.extend(sorted(found))
        elif path.is_file():
            files.append(path)
        else:
            raise ValueError(f"找不到输入文件: {path}")
    return files


def convert(inputs: Iterable[Path]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for path in iter_input_files(inputs):
        for record in load_records(path):
            rows.append(record_to_row(record))
    return rows


def default_output(inputs: list[Path]) -> Path:
    files = iter_input_files(inputs)
    if len(files) == 1:
        return files[0].with_suffix(".xlsx")
    return Path("llm_log.xlsx")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="把语音助手 LLM 日志（JSON / JSONL）转成 Excel。"
    )
    parser.add_argument(
        "inputs",
        nargs="+",
        type=Path,
        help="JSON、JSONL 文件，或包含这些文件的目录",
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        help="输出 xlsx 路径。默认与单个输入文件同名，多个输入时为 ./llm_log.xlsx",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        rows = convert(args.inputs)
        output = args.output or default_output(args.inputs)
        if not rows:
            print("没有读到任何记录", file=sys.stderr)
            return 1
        write_excel(rows, output)
    except (OSError, ValueError) as exc:
        print(exc, file=sys.stderr)
        return 1

    judged = {}
    for row in rows:
        name = row.get("jev判断模型") or "(空)"
        judged[name] = judged.get(name, 0) + 1
    summary = ", ".join(f"{name} {count}" for name, count in sorted(judged.items()))
    print(f"已写入 {len(rows)} 行 -> {output}")
    print(f"jev判断模型: {summary}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
